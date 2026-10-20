#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT
parser = argparse.ArgumentParser()
parser.add_argument("--source", default=str(BASE / "MANUSCRIPT_JAE_V1_2.md"))
parser.add_argument("--output", default=str(ROOT / "build" / "FROG_JAE_ANON_MAIN_V1_2.docx"))
args = parser.parse_args()
SRC = Path(args.source)
OUT = Path(args.output)
OUT.parent.mkdir(parents=True, exist_ok=True)


def add_inline(paragraph, text):
    # Render Markdown inline-code markers as ordinary manuscript text.
    text = text.replace("`", "")
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def set_run_font(run, name="Times New Roman", size=12):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), name)


def format_paragraph(paragraph, *, first_line=None, hanging=None, keep=False):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_line is not None:
        fmt.first_line_indent = Inches(first_line)
    if hanging is not None:
        fmt.left_indent = Inches(hanging)
        fmt.first_line_indent = Inches(-hanging)
    fmt.keep_with_next = keep
    for run in paragraph.runs:
        set_run_font(run)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    visible = OxmlElement("w:t")
    visible.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, visible, fld_end])
    set_run_font(run, size=10)


def is_markdown_table_divider(line):
    """A markdown alignment row, e.g. |:---|---:|:---:|."""
    if not line.strip().startswith("|"):
        return False
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells)


def table_cells(raw):
    """The frozen SI uses ordinary pipe tables, not escaped interior pipes."""
    if not raw.strip().startswith("|"):
        raise ValueError("not a Markdown table row")
    return [cell.strip() for cell in raw.strip().strip("|").split("|")]


def add_markdown_table(doc, raw_lines, *, page_width_inches=6.27):
    """Convert the SI's markdown grids into real, readable Word tables.

    Fail closed rather than silently print a pipe-delimited row as a paragraph:
    that defect concealed all fourteen SI tables during the first build.
    """
    if len(raw_lines) < 3 or not is_markdown_table_divider(raw_lines[1]):
        raise ValueError("Markdown table is missing a header alignment rule")
    rows = [table_cells(raw) for raw in raw_lines]
    ncols = len(rows[0])
    if ncols < 2 or ncols > 8 or any(len(r) != ncols for r in rows):
        raise ValueError("Malformed Markdown table: column count not stable")
    if any(is_markdown_table_divider(row) for row in raw_lines[2:]):
        raise ValueError("Unexpected table divider in body")
    rows.pop(1)  # Alignment syntax is not an actual Word row.

    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Reserve a minimum width for numeric columns while assigning additional
    # width to long-description columns. The 6.27in area is A4 minus margins.
    maxlens = [max(len(row[i]) for row in rows) for i in range(ncols)]
    weights = [max(8.0, min(65.0, float(x))) for x in maxlens]
    minimum = 0.63 if ncols >= 5 else 0.77
    remaining = page_width_inches - minimum * ncols
    if remaining <= 0:
        raise ValueError("too many columns for document page width")
    widths = [minimum + remaining * w / sum(weights) for w in weights]
    font_size = 8.4 if ncols >= 5 else 9.3

    def fill_row(row, values, header):
        row_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)
        if header:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_pr.append(repeat)
        for i, text in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, text)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.keep_with_next = header
            for run in paragraph.runs:
                set_run_font(run, size=font_size)
                if header:
                    run.bold = True
            if header:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEF2F6")
                cell._tc.get_or_add_tcPr().append(shading)

    fill_row(table.rows[0], rows[0], True)
    for values in rows[1:]:
        fill_row(table.add_row(), values, False)
    return table


def build():
    source = SRC.read_text(encoding="utf-8")
    title_line = next((x[2:].strip() for x in source.splitlines() if x.startswith("# ")), "")
    if not title_line:
        raise SystemExit("manuscript title heading not found")
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Inches(1))

    # Continuous line numbering for double-anonymized review.
    sect_pr = section._sectPr
    line_num = sect_pr.find(qn("w:lnNumType"))
    if line_num is None:
        line_num = OxmlElement("w:lnNumType")
        sect_pr.append(line_num)
    line_num.set(qn("w:countBy"), "1")
    line_num.set(qn("w:start"), "1")
    line_num.set(qn("w:restart"), "continuous")
    line_num.set(qn("w:distance"), "360")

    add_page_field(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    mode = "body"
    raw_lines = source.splitlines()
    position = 0
    expected_tables = 0
    while position < len(raw_lines):
        raw_line = raw_lines[position]
        position += 1
        line = raw_line.rstrip()
        if not line:
            continue
        if line.strip() == "---":
            continue
        if line.startswith("> "):
            line = line[2:].lstrip()

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            r.bold = True
            set_run_font(r, size=14)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.space_after = Pt(12)
            continue

        if line.startswith("## "):
            heading = line[3:]
            p = doc.add_paragraph()
            r = p.add_run(heading)
            r.bold = True
            set_run_font(r)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            mode = "references" if heading == "References" else (
                "abstract" if heading == "Abstract" else "body"
            )
            continue

        if line.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(line[4:])
            r.bold = True
            set_run_font(r)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            continue

        if mode == "abstract" and re.match(r"^[1-5]\. ", line):
            p = doc.add_paragraph()
            add_inline(p, line)
            format_paragraph(p, first_line=0)
            continue

        if mode == "references":
            p = doc.add_paragraph()
            add_inline(p, line)
            format_paragraph(p, hanging=0.35)
            continue

        if line.startswith("|"):
            block = [line]
            while position < len(raw_lines) and raw_lines[position].lstrip().startswith("|"):
                block.append(raw_lines[position].rstrip())
                position += 1
            add_markdown_table(doc, block)
            expected_tables += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, line)
        first_line = 0.3
        if line.startswith("acoustic community;") or line.startswith("**Figure ") or line.startswith("Figure ") or line.startswith("|"):
            first_line = 0
        format_paragraph(p, first_line=first_line)


    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.size is None:
                set_run_font(run)

    if len(doc.tables) != expected_tables or any(
        p.text.lstrip().startswith("|") for p in doc.paragraphs
    ):
        raise ValueError("Markdown table text leaked into Word document")
    if SRC.name == "supporting_information.md" and expected_tables != 14:
        raise ValueError(
            f"Supporting Information table count drift: {expected_tables}, expected 14"
        )

    props = doc.core_properties
    props.title = title_line
    props.author = ""
    props.last_modified_by = ""
    props.subject = "Journal of Animal Ecology anonymous main manuscript"
    props.comments = ""
    props.keywords = ""

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
