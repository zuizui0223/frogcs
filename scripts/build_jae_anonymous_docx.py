#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT
parser = argparse.ArgumentParser()
parser.add_argument("--source", default=str(BASE / "MANUSCRIPT_JAE_V0_4.md"))
parser.add_argument("--output", default=str(ROOT / "build" / "FROG_JAE_ANON_MAIN_V0_4.docx"))
args = parser.parse_args()
SRC = Path(args.source)
OUT = Path(args.output)
OUT.parent.mkdir(parents=True, exist_ok=True)


def add_inline(paragraph, text):
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def set_run_font(run, name="Times New Roman", size=12):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), name)


def format_paragraph(paragraph, *, first_line=None, hanging=None, keep=False):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_line is not None:
        fmt.first_line_indent = Inches(first_line)
    if hanging is not None:
        fmt.left_indent = Inches(hanging)
        fmt.first_line_indent = Inches(-hanging)
    fmt.keep_with_next = keep
    for run in paragraph.runs:
        set_run_font(run)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    visible = OxmlElement("w:t")
    visible.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, visible, fld_end])
    set_run_font(run, size=10)


def build():
    source = SRC.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Inches(1))

    # Continuous line numbering for double-anonymized review.
    sect_pr = section._sectPr
    line_num = sect_pr.find(qn("w:lnNumType"))
    if line_num is None:
        line_num = OxmlElement("w:lnNumType")
        sect_pr.append(line_num)
    line_num.set(qn("w:countBy"), "1")
    line_num.set(qn("w:start"), "1")
    line_num.set(qn("w:restart"), "continuous")
    line_num.set(qn("w:distance"), "360")

    add_page_field(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    mode = "body"
    for raw_line in source.splitlines():
        line = raw_line.rstrip()
        if not line:
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            r.bold = True
            set_run_font(r, size=14)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.space_after = Pt(12)
            continue

        if line.startswith("## "):
            heading = line[3:]
            p = doc.add_paragraph()
            r = p.add_run(heading)
            r.bold = True
            set_run_font(r)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            mode = "references" if heading == "References" else (
                "abstract" if heading == "Abstract" else "body"
            )
            continue

        if line.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(line[4:])
            r.bold = True
            set_run_font(r)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            continue

        if mode == "abstract" and re.match(r"^[1-5]\. ", line):
            p = doc.add_paragraph()
            add_inline(p, line)
            format_paragraph(p, first_line=0)
            continue

        if mode == "references":
            p = doc.add_paragraph()
            add_inline(p, line)
            format_paragraph(p, hanging=0.35)
            continue

        p = doc.add_paragraph()
        add_inline(p, line)
        first_line = 0.3
        if line.startswith("acoustic community;") or line.startswith("**Figure ") or line.startswith("Figure "):
            first_line = 0
        format_paragraph(p, first_line=first_line)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.size is None:
                set_run_font(run)

    props = doc.core_properties
    props.title = "Recent rainfall predicts broader frog acoustic participation without stronger residual co-calling associations"
    props.author = ""
    props.last_modified_by = ""
    props.subject = "Journal of Animal Ecology anonymous main manuscript"
    props.comments = ""
    props.keywords = ""

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
