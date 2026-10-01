#!/usr/bin/env python3
"""Inspect WFTS file structure without browsing biological response values.

Supported:
- CSV/TSV: reads only the header row and file byte size.
- XLSX: reads workbook metadata and the first row of each worksheet in read-only mode.
- JSON: reports only top-level type/keys; it does not enumerate records.

This utility is for Stage 2 of WFTS_RECEIPT_AND_PREFLIGHT_PROTOCOL_V0_1.md.
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path


def inspect_delimited(path: Path, delimiter: str) -> dict:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        try:
            header = next(reader)
        except StopIteration:
            header = []
    return {
        "format": "csv" if delimiter == "," else "tsv",
        "columns": header,
        "column_count": len(header),
        "rows_read": 1 if header else 0,
        "response_rows_read": 0,
    }


def inspect_xlsx(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("XLSX inspection requires openpyxl") from exc
    wb = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    try:
        for ws in wb.worksheets:
            row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            columns = ["" if x is None else str(x) for x in row]
            sheets.append(
                {
                    "sheet": ws.title,
                    "columns": columns,
                    "column_count": len(columns),
                    "rows_read": 1 if columns else 0,
                    "response_rows_read": 0,
                }
            )
    finally:
        wb.close()
    return {
        "format": "xlsx",
        "sheet_count": len(sheets),
        "sheets": sheets,
        "response_rows_read": 0,
    }


def inspect_json(path: Path) -> dict:
    # JSON cannot generally expose keys without parsing the document, so only a
    # small structural object is supported. Do not use this on a large response
    # record array before the adapter is frozen.
    obj = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(obj, dict):
        return {
            "format": "json",
            "top_level_type": "object",
            "top_level_keys": sorted(obj.keys()),
            "response_rows_read": 0,
            "warning": "Only top-level object keys are reported.",
        }
    return {
        "format": "json",
        "top_level_type": type(obj).__name__,
        "top_level_keys": [],
        "response_rows_read": 0,
        "warning": (
            "Top-level non-object JSON may represent response records; do not inspect "
            "further before a format-specific adapter is frozen."
        ),
    }


def inspect(path: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        detail = inspect_delimited(path, ",")
    elif suffix in {".tsv", ".txt"}:
        detail = inspect_delimited(path, "\t")
    elif suffix == ".xlsx":
        detail = inspect_xlsx(path)
    elif suffix == ".json":
        detail = inspect_json(path)
    else:
        detail = {
            "format": suffix.lstrip(".") or "unknown",
            "response_rows_read": 0,
            "warning": "Unsupported container: byte receipt only; do not browse manually.",
        }
    return {
        "filename": path.name,
        "byte_size": path.stat().st_size,
        **detail,
    }


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("files", nargs="+", type=Path)
    p.add_argument("--output", required=True, type=Path)
    args = p.parse_args()

    results = []
    for path in args.files:
        if not path.is_file():
            raise SystemExit(f"not a regular file: {path}")
        results.append(inspect(path))

    receipt = {
        "receipt_type": "wfts_schema_only_v0_1",
        "files": results,
        "file_count": len(results),
        "response_values_interpreted": False,
        "response_rows_read": 0,
        "note": (
            "Schema-only inspection. Column names/sheet names may be used to write a "
            "purely syntactic adapter; biological response rows were not browsed."
        ),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    print(
        {
            "status": "PASS",
            "files": len(results),
            "response_values_interpreted": False,
            "response_rows_read": 0,
        }
    )


if __name__ == "__main__":
    main()
