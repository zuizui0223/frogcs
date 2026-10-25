#!/usr/bin/env python3
"""Public original-data schema only: direct risk/cue experimental programmes.

Inspect source SHA, documented treatment keys, rows and NULLs. DO NOT emit
outcomes, individual source records, personal IDs or third-party raw files.
Downloaded bytes are temporary within the GitHub Actions runner.
"""
import argparse
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import time
import urllib.error
import urllib.request
import zipfile
from collections import Counter

SOURCES = (
    ("urushihara_predation", 4721260, "predation.trials.2024.csv", "10.5061/dryad.31zcrjdzw", 1958),
    ("urushihara_early_behavior", 4721257, "behavior.step.1.2024.csv", "10.5061/dryad.31zcrjdzw", 3054),
    ("urushihara_late_traits", 4721258, "morphology.behavior.life-history.traits.step.2.2024.csv", "10.5061/dryad.31zcrjdzw", 8691),
    ("urushihara_early_morphology", 4721259, "morphology.step.1.2024.csv", "10.5061/dryad.31zcrjdzw", 1997),
    ("thapa_predator_odors", 2986323, "predatordata.xlsx", "10.5061/dryad.prr4xgxtr", 29052),
    ("baker_multitrophic", 4050103, "PredationAltersCommunityStructure2024_Data.csv", "10.5061/dryad.6q573n68k", 37671),
    ("baker_tadpoles", 4050104, "PredationAltersCommunityStructure2024_DataTadpoles.csv", "10.5061/dryad.6q573n68k", 9475)
)

def download(fid, filename):
    urls = (f"https://datadryad.org/downloads/file_stream/{fid}",
            f"https://datadryad.org/api/v2/files/{fid}/download")
    if fid == 2986323:
        # NLM official 2026 PMC AWS article-version object, CC-BY supplement.
        # Independent publication-source fallback; not asserted byte-identical to Dryad.
        urls += ("https://pmc-oa-opendata.s3.amazonaws.com/"
                 "PMC10965031.1/arae014_suppl_supplementary_data.xlsx",)
    errors = []
    for url in urls:
        try:
            req=urllib.request.Request(url,headers={
                "User-Agent": "Mozilla/5.0 (compatible; Scientific-Replication-Data-Audit/1.0)",
                "Accept": "application/octet-stream,text/csv,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
                "Referer": "https://datadryad.org/"})
            with urllib.request.urlopen(req, timeout=40) as resp:
                blob=resp.read(1_000_000)
                status=resp.status
                finalurl=resp.url
                ctype=resp.headers.get("Content-Type","")
            if status!=200 or len(blob)<50:
                errors.append({"endpoint":url,"reason":"bad_status_or_empty","status":status})
                continue
            if blob[:5].lower()==b"<html" or b"<html" in blob[:500].lower():
                errors.append({"endpoint":url,"reason":"html_instead_of_original_bytes","status":status})
                continue
            if filename.endswith(".xlsx") and not zipfile.is_zipfile(io.BytesIO(blob)):
                errors.append({"endpoint":url,"reason":"not_xlsx_zip","status":status})
                continue
            if filename.endswith(".csv") and (b"\n" not in blob or b"," not in blob[:1000]):
                errors.append({"endpoint":url,"reason":"not_csv","status":status})
                continue
            return blob, {"endpoint":url,"resolved_host":urllib.parse.urlparse(finalurl).netloc,
                          "reported_mime":ctype}, errors
        except (urllib.error.URLError, ValueError, TimeoutError, OSError) as ex:
            status=getattr(ex,"code",None)
            errors.append({"endpoint":url,"reason":type(ex).__name__,"status":status,
                           "message":str(ex)[:140]})
    return None,None,errors

def table_schema(data):
    text = data.decode("utf-8-sig")
    reader=csv.DictReader(io.StringIO(text))
    fields=reader.fieldnames or []
    rows=list(reader)
    report={"format":"csv","n_rows":len(rows),"fields":fields}
    if len(rows)>200_000:
        raise ValueError("oversized schema source")
    # Basic column nulls, but not outcome distributions.
    report["null_counts"]={f:sum(not str(r.get(f) or "").strip() for r in rows) for f in fields}
    # Only source-defined grouping fields, not IDs or outcome data.
    tags=[f for f in fields if f.lower() in {
          "treatment","treatment.1","treatment.2","trial","cue","condition","group","treat","stage",
          "type","size.treatment"}]
    report["observed_groups"]={f:dict(Counter((r.get(f) or "<missing>").strip() for r in rows)) for f in tags}
    return report

def xlsx_schema(data):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        itr=ws.iter_rows(values_only=True)
        hdr=list(next(itr))
        count=0;missing=Counter()
        groups={}
        pails=set();clutches=set();clutch_pails=set()
        fields=[str(v).strip() if v is not None else "" for v in hdr]
        for row in itr:
            if all(x is None for x in row):
                continue
            count+=1
            raw={fields[i]:v for i,v in enumerate(row) if i<len(fields)}
            pails.add(str(raw.get("Pail")))
            clutches.add(str(raw.get("Clutch")))
            clutch_pails.add((str(raw.get("Clutch")),str(raw.get("Pail"))))
            for i,f in enumerate(fields):
                if i>=len(row) or row[i] is None:
                    missing[f]+=1
                if f.lower() in ("treatment","cue","tx"):
                    groups.setdefault(f,Counter())[str(row[i]) if i<len(row) else "null"]+=1
        sheets.append({"sheet":ws.title,"header":fields,"n_rows":count,
                       "n_distinct_pails":len(pails),"n_distinct_clutches":len(clutches),
                       "n_distinct_clutch_pails":len(clutch_pails),
                       "null_counts":dict(missing),"group_columns":{k:dict(v) for k,v in groups.items()}})
    return {"format":"xlsx","sheets":sheets}

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--out",required=True)
    p=parser.parse_args()
    records=[]
    for label,file_id,filename,doi,expected_size in SOURCES:
        raw,source,errors=download(file_id,filename)
        result={"source_id":label,"doi":doi,"file_id":file_id,"filename":filename,
                "catalog_size_bytes":expected_size}
        if raw is None:
            result.update({"status":"SOURCE_DOWNLOAD_UNAVAILABLE","attempts":errors})
            print("SOURCE_BLOCKED",label, json.dumps(errors))
        else:
            result.update({"status":"ORIGINAL_BYTES_READ","download":source,
                           "byte_count":len(raw),"sha256":hashlib.sha256(raw).hexdigest(),
                           "size_matches_catalog":len(raw)==expected_size})
            try:
                schema=xlsx_schema(raw) if filename.endswith(".xlsx") else table_schema(raw)
                result.update({"status":"SOURCE_SCHEMA_VERIFIED","schema":schema})
                print("SOURCE_SCHEMA_VERIFIED",label,len(raw),
                      "rows",schema.get("n_rows",sum(x["n_rows"] for x in schema.get("sheets",[]))))
                print("SANITIZED_SOURCE_SCHEMA_JSON",label,json.dumps({"sha256":result["sha256"],"schema":schema},ensure_ascii=False))
            except Exception as ex:
                result["status"]="SOURCE_PARSE_ERROR"
                result["parse_error"]=type(ex).__name__+": "+str(ex)[:180]
                print("SOURCE_PARSE_ERROR",label,result["parse_error"])
        records.append(result)
    out={"programme":"ONTO-TIME original direct-risk source gate v0.1",
         "data_origin":"public published Dryad files",
         "scientific_status":"SCHEMA_ONLY_NO_OUTCOME_REANALYSIS",
         "raw_bytes_redistributed":False,
         "n_sources":len(records),
         "n_source_schema_verified":sum(r["status"]=="SOURCE_SCHEMA_VERIFIED" for r in records),
         "records":records}
    Path(p.out).write_text(json.dumps(out,ensure_ascii=False,indent=2)+"\n")
    print("SCHEMA_GATE_SUMMARY",json.dumps({"valid":out["n_source_schema_verified"],
                                           "attempted":out["n_sources"]}))
    # Fail closed on all sources blocked; partial success documented without pretending completeness.
    if not out["n_source_schema_verified"]:
        raise SystemExit("NO_ORIGINAL_DATA: all public download endpoints unavailable")


if __name__=="__main__":
    main()
