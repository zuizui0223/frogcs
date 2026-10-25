#!/usr/bin/env python3
"""ONTO-TIME: original Thapa et al. 2024 predator metamorphosis cue-transfer test.

Public source: NLM PMC AWS (CC-BY), PMC10965031 version 1.
Frozen output is grouped/pail-level source-derived estimates only: no raw xlsx,
individual movement counts, sensitive collection data, or pail IDs emitted.

This is an exploratory *independent source reanalysis* within a multi-system
programme. Odor training is NOT a frog-call acoustic alarm experiment.
"""
import argparse
from collections import Counter, defaultdict
import hashlib
import io
import json
import math
from pathlib import Path
import random
import urllib.request

SRC="https://pmc-oa-opendata.s3.amazonaws.com/PMC10965031.1/arae014_suppl_supplementary_data.xlsx"
SOURCE_SHA256="a200b66072c5de204efe91536066d66c9cf8b513c5075c321fc44391544d9fa5"
DOCUMENTED_2024_TESTED={"beetle":360,"salamander":357}
ACTUAL_ROWS={"beetle":308,"salamander":280}
SEED=20261025
B=10000
TRAIN=("sham","larval","adult")
TEST=("larval","adult")

def stage(value):
    s=str(value).strip().casefold()
    if "water" in s:
        return "sham"
    if "larval" in s or "larvae" in s:
        return "larval"
    if "adult" in s:
        return "adult"
    raise ValueError("unrecognized source cue/treatment: "+repr(value))

def mean(v):
    return sum(v)/len(v) if v else None

def percentile(vals, q):
    x=sorted(vals); j=(len(x)-1)*q; lo=int(j);hi=min(lo+1,len(x)-1)
    return x[lo]*(hi-j)+x[hi]*(j-lo)

def rounded(x,digits=6):
    return None if x is None else round(x,digits)

def fetch():
    request=urllib.request.Request(SRC,headers={"User-Agent":"ONTO-TIME-Original-Cue-Study/0.1"})
    with urllib.request.urlopen(request,timeout=35) as res:
        raw=res.read(1_000_000)
    sha=hashlib.sha256(raw).hexdigest()
    if sha!=SOURCE_SHA256:
        raise ValueError("source version drift: original PMC AWS XLSX SHA does not match fixed contract "+sha)
    return raw

def source_rows(wb, name):
    ws=wb["Beetle" if name=="beetle" else "Salamader"]
    itr=ws.iter_rows(values_only=True)
    fields=[str(x).strip() for x in next(itr)]
    expected=["Pail","Clutch","Treatment","Cue","Pre","Post","Change"] if name=="beetle" else [
        "Pail","Clutch","Tx","Cue","Pre","Post","Change"]
    if fields[:7]!=expected:
        raise ValueError("unexpected source headers "+str(fields))
    pails={}
    rawgroups=Counter()
    pail_clutches=defaultdict(set)
    count=0
    pre_not_active=0
    for row in itr:
        if all(x is None for x in row[:7]):
            continue
        if len(row)<7 or any(v is None for v in row[:7]):
            raise ValueError("source has incomplete activity row")
        pid=str(row[0]).strip()
        clutch=str(row[1]).strip()
        training=stage(row[2]); test=stage(row[3])
        if test=="sham":
            raise ValueError("sham test odor incorrectly entered")
        pre,post,change=(float(row[i]) for i in (4,5,6))
        if not all(math.isfinite(v) for v in (pre,post,change)):
            raise ValueError("nonfinite activity row")
        if abs((post-pre)-change)>1e-7:
            raise ValueError("source Change is not Post minus Pre")
        if pre<6:
            pre_not_active+=1
        if pid not in pails:
            pails[pid]={"training":training,"values":{"adult":[],"larval":[]},"clutches":set()}
        p=pails[pid]
        if p["training"]!=training:
            raise ValueError("randomization unit Pail has multiple conditioning treatments")
        p["values"][test].append(change)
        p["clutches"].add(clutch)
        pail_clutches[pid].add(clutch)
        rawgroups[(training,test)]+=1
        count+=1
    if count!=ACTUAL_ROWS[name]:
        raise ValueError("observed XLSX rows changed or previous status stale")
    if len(pails)!=18:
        raise ValueError("source pail count changed")
    bytrain=Counter(p["training"] for p in pails.values())
    if bytrain!={"sham":6,"larval":6,"adult":6}:
        raise ValueError("source randomization pail counts are not six per arm: "+str(bytrain))
    empty=[p for p in pails.values() if not p["values"]["adult"] or not p["values"]["larval"]]
    if empty:
        print("PAIL_TEST_CUE_SCHEMA_GATE",json.dumps({
           "system":name,
           "n_pails":len(pails),
           "n_pails_with_both_cues":len(pails)-len(empty),
           "incomplete_by_training":dict(Counter(p["training"] for p in empty)),
           "missing_cue_by_training":dict(Counter(p["training"]+"|"+c for p in empty
                                                   for c in TEST if not p["values"][c]))
        }))
        raise ValueError("missing test-cue within experimental pail; no silent sample filtering")
    grouped={t:[] for t in TRAIN}
    for p in pails.values():
        grouped[p["training"]].append({v:mean(p["values"][v]) for v in TEST})
    return {
        "n_data_rows":count,"n_pails":len(pails),"n_distinct_clutches":
            len(set().union(*(p["clutches"] for p in pails.values()))),
        "n_multiclutch_pails":sum(len(v)>1 for v in pail_clutches.values()),
        "n_clutch_pail_cells":sum(len(v) for v in pail_clutches.values()),
        "pre_below_6":pre_not_active,"source_train_cue_counts":
            {t:{c:rawgroups[(t,c)] for c in TEST} for t in TRAIN},
        "pail_counts_by_training":dict(bytrain),
        "grouped":grouped
    }

def effects(grouped):
    mu={t:{cue:mean([p[cue] for p in grouped[t]]) for cue in TEST} for t in TRAIN}
    cross_larval_to_adult=mu["larval"]["adult"]-mu["sham"]["adult"]
    cross_adult_to_larval=mu["adult"]["larval"]-mu["sham"]["larval"]
    same_larval=mu["larval"]["larval"]-mu["sham"]["larval"]
    same_adult=mu["adult"]["adult"]-mu["sham"]["adult"]
    cross=(cross_larval_to_adult+cross_adult_to_larval)/2
    same=(same_larval+same_adult)/2
    return {"cross_larval_training_to_adult_cue":cross_larval_to_adult,
            "cross_adult_training_to_larval_cue":cross_adult_to_larval,
            "same_larval_training_larval_cue":same_larval,
            "same_adult_training_adult_cue":same_adult,
            "cross_stage_mean":cross,"within_stage_mean":same,
            "cross_minus_within_stage":cross-same}, mu

def bootstrap(grouped):
    rng=random.Random(SEED)
    bs={x:[] for x in effects(grouped)[0]}
    for _ in range(B):
        sub={}
        for key in TRAIN:
            p=grouped[key];sub[key]=[p[rng.randrange(len(p))] for __ in p]
        effects1,_=effects(sub)
        for key,val in effects1.items():
            bs[key].append(val)
    point,_=effects(grouped)
    return {key:{"estimate":rounded(point[key]),
                 "bootstrap95":(rounded(percentile(v,.025)),rounded(percentile(v,.975))),
                 "n_training_pails":{t:len(grouped[t]) for t in TRAIN}}
            for key,v in bs.items()}

def influence(grouped):
    output={}
    full,_=effects(grouped)
    for key in ("cross_stage_mean","within_stage_mean","cross_minus_within_stage"):
        vals=[]
        for t in TRAIN:
            for i in range(len(grouped[t])):
                short={n:(arr[:i]+arr[i+1:] if n==t else arr)
                       for n,arr in grouped.items()}
                val,_=effects(short)
                vals.append(val[key])
        output[key]={"min_leave_one_pail":rounded(min(vals)),
                     "max_leave_one_pail":rounded(max(vals)),
                     "sign_flips":sum((x<0)!=(full[key]<0) for x in vals),
                     "n_leave_one_pail_analyses":len(vals)}
    return output

def analyze():
    from openpyxl import load_workbook
    raw=fetch()
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    if set(wb.sheetnames)!={"Beetle","Salamader"}:
        raise ValueError("source workbook sheets unexpected")
    out={"status":"ORIGINAL_SOURCE_DIRECT_PAIL_LEVEL_RISK_CUE_REANALYSIS",
         "original_publication":"10.1093/beheco/arae014",
         "official_source_url":SRC,
         "origin_pmc":"PMC10965031.1",
         "sha256":SOURCE_SHA256,
         "claim_boundary":"Measured chemical odor learning/generalization, NOT acoustic warning or direct reciprocal predation",
         "bootstrap":{"method":"equal conditioning Pail with 6 pails in each train arm","draws":B,
                      "seed":SEED, "uncertainty":"descriptive percentile bootstrap, not independent external replication"},
         "systems":{}}
    for name in ("beetle","salamander"):
        source=source_rows(wb,name)
        stats=bootstrap(source["grouped"])
        point,mu=effects(source["grouped"])
        out["systems"][name]={
            "n_source_rows":source["n_data_rows"],"n_pails":source["n_pails"],
            "n_clutches":source["n_distinct_clutches"],
            "n_clutch_pail_cells":source["n_clutch_pail_cells"],
            "n_multiclutch_pails":source["n_multiclutch_pails"],
            "n_pre_lt_6":source["pre_below_6"],
            "n_paper_reported_tested":DOCUMENTED_2024_TESTED[name],
            "source_minus_published_reported_n":source["n_data_rows"]-DOCUMENTED_2024_TESTED[name],
            "group_cell_source_counts":source["source_train_cue_counts"],
            "pail_level_mean_change":{t:{cue:rounded(mu[t][cue]) for cue in TEST} for t in TRAIN},
            "contrasts":stats,
            "leave_one_pail_out":influence(source["grouped"]),
            "within_study_sampling_limitation":"Original source row count differs from publication reported tested n; selection/omission mechanism not identified."
        }
    return out

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--out",required=True)
    a=p.parse_args()
    result=analyze()
    Path(a.out).write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    # Exposes aggregated pail-level coefficients only, never individual values or IDs.
    print("THAPA_ORIGINAL_PAIL_RESULTS_BEGIN")
    print(json.dumps(result,ensure_ascii=False,indent=2))
    print("THAPA_ORIGINAL_PAIL_RESULTS_END")

if __name__=="__main__":
    main()
