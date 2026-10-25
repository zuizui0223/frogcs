#!/usr/bin/env python3
"""Response-blind Dryad Carroll et al. (2023) workbook schema audit.

The fixed source is Dryad DOI 10.25338/B8G92P, version 218704,
funneltrapdata.xlsx file id 2077176 and README file id 2077175.
Never reinterpret unlabeled sheets, columns, numeric values or missing rows
as biological zero, abundance, a directly aged overwintered cohort, or
trophic attack events. This is a metadata-only source preflight, not H1.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import sys
import tempfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import openpyxl

DOI = "10.25338/B8G92P"
VERSION = 218704
WORKBOOK_ID = 2077176
WORKBOOK_SHA256 = "302f1d976725ff926c682331a34e302262adbdb3be1f80fc67c1ccc990343f6f"
WORKBOOK_BYTES = 196918
README_ID = 2077175
README_SHA256 = "d2729af5cf0e685b15a26fc80f9ce589010c67e3cda545c9b2160c6b8906a5ff"
ROOT = Path(__file__).resolve().parents[2]
URLS = {
    "workbook": [
        "https://datadryad.org/api/v2/files/2077176/download",
        "https://datadryad.org/downloads/file_stream/2077176",
    ],
    "readme": [
        "https://datadryad.org/api/v2/files/2077175/download",
        "https://datadryad.org/downloads/file_stream/2077175",
    ],
}

def safe_download(kind, output, expect_sha, max_bytes):
    diagnostics = []
    for url in URLS[kind]:
        req = Request(url, headers={
            "User-Agent": "frogcs-stage-legacy-source-audit/1.0 (reproducible ecology; github.com/zuizui0223/frogcs)",
            "Accept": "application/octet-stream, */*",
        })
        try:
            with urlopen(req, timeout=50) as response:
                raw = response.read(max_bytes + 1)
            if len(raw) > max_bytes:
                raise ValueError("FILE_OVERSIZE")
            got = hashlib.sha256(raw).hexdigest()
            if got != expect_sha:
                diagnostics.append({"url":url,"status":"digest_mismatch", "size":len(raw), "sha256":got})
                continue
            output.write_bytes(raw)
            return {"status":"PINNED_DIGEST_OK", "source_url":url,
                    "bytes":len(raw), "sha256":got, "failed_attempts": diagnostics}
        except (HTTPError, URLError, TimeoutError, OSError, ValueError) as exc:
            diagnostics.append({"url":url, "status":type(exc).__name__, "message":str(exc)[:200]})
    return {"status":"BLOCKED", "attempts":diagnostics}

def string_cells(row):
    """Only header labels, never raw numeric biological outcome values."""
    return [str(v).strip() for v in row if isinstance(v,str) and v.strip()]

def cell_pattern(row):
    out = Counter()
    for v in row:
        if v is None: continue
        if isinstance(v, (date, datetime)): out["dates"] += 1
        elif isinstance(v, (int,float)): out["numbers"] += 1
        elif isinstance(v,str): out["strings"] += 1
        else: out["other"] += 1
    return dict(out)

def get_schema(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets=[]
    for ws in book.worksheets:
        probe=list(ws.iter_rows(min_row=1,max_row=min(ws.max_row,6),values_only=True))
        header_candidates=[]
        for index, row in enumerate(probe,1):
            labels=string_cells(row)
            if len(labels)>=2 and all(len(x)<=95 for x in labels):
                header_candidates.append({"row":index, "labels":labels[:65]})
        sheets.append({
            "sheet_name":ws.title, "max_row":ws.max_row,
            "max_column":ws.max_column,
            "top_rows_type_pattern":[{"row":i,"types":cell_pattern(row)}
                                      for i,row in enumerate(probe,1)],
            "candidate_column_labels":header_candidates[:4],
            "outcome_cells_disclosed":False,
        })
    book.close()
    return sheets

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, help="locally supplied pinned Dryad workbook; if omitted download from Dryad")
    ap.add_argument("--readme", type=Path,help="locally supplied pinned README")
    ap.add_argument("--out",type=Path)
    args=ap.parse_args()
    result={
        "programme":"frogcs stage legacy G0","source":"Carroll et al. (2023)",
        "doi":DOI,"dryad_version_id":VERSION,
        "source_workbook_file_id":WORKBOOK_ID,
        "source_workbook_sha256_expected":WORKBOOK_SHA256,
        "source_workbook_bytes_expected":WORKBOOK_BYTES,
        "schema_blind_to_numeric_biological_outcomes":True,
        "scope":"file integrity; sheet names/dimensions/headers only; no adult/larval effect estimates",
    }
    with tempfile.TemporaryDirectory() as d:
        p=Path(d)
        paths={}
        for kind,supplied,sha,size in [
            ("workbook",args.xlsx,WORKBOOK_SHA256,500000),
            ("readme",args.readme,README_SHA256,20000),
        ]:
            target=supplied or p/("funneltrapdata.xlsx" if kind=="workbook" else "README.md")
            if supplied:
                got=hashlib.sha256(target.read_bytes()).hexdigest()
                access={"status":"PINNED_DIGEST_OK" if got==sha else "DIGEST_MISMATCH",
                        "sha256":got,"bytes":target.stat().st_size,"source_url":"user_supplied_pinned_copy"}
            else:
                access=safe_download(kind,target,sha,size)
            result[kind+"_access"]=access
            if access["status"]=="PINNED_DIGEST_OK":paths[kind]=target
        if "workbook" in paths:
            assert paths["workbook"].stat().st_size==WORKBOOK_BYTES,"Original Dryad size changed"
            result["sheets"]=get_schema(paths["workbook"])
        else:
            result["sheets"]=[]
        if "readme" in paths:
            text=paths["readme"].read_text(encoding="utf-8-sig",errors="replace")
            result["readme_line_count"]=len(text.splitlines())
            result["readme_text"]=text[:8000]
        result["G0_status"]=("METADATA_AUDIT_COMPLETED_OUTCOME_BLIND"
                             if "workbook" in paths and "readme" in paths
                             else "SOURCE_BYTES_NOT_AVAILABLE_NO_ANALYSIS")
        result["lagged_predator_proxy_validated"]=False
        result["predation_inferred"]=False
    if args.out:
        args.out.parent.mkdir(parents=True,exist_ok=True)
        args.out.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print("BEGIN_CARROLL_SOURCE_G0_JSON")
    print(json.dumps(result,indent=2,ensure_ascii=False))
    print("END_CARROLL_SOURCE_G0_JSON")
    if result["G0_status"]!="METADATA_AUDIT_COMPLETED_OUTCOME_BLIND":
        sys.exit(4)

if __name__=="__main__": main()
